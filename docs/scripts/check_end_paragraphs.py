import docx

doc = docx.Document('c:/Users/prem/OSINT-breach-Finder-main/docs/academic/BreachShield_Progress_Report_3.docx')
print('Total paragraphs:', len(doc.paragraphs))
for i in range(45, len(doc.paragraphs)):
    t = doc.paragraphs[i].text.strip()
    if t:
        print(f'P{i}: \"{t}\"')
    else:
        # check if it has sectPr
        xml = doc.paragraphs[i]._element.xml
        if 'sectPr' in xml:
            print(f'P{i}: [EMPTY WITH sectPr]')
