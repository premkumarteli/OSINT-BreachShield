import docx
from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import shutil

template_path = 'c:/Users/prem/OSINT-breach-Finder-main/docs/academic/Netvisor_Project Work Progress Report 1.docx'
academic_dir = 'c:/Users/prem/OSINT-breach-Finder-main/docs/academic'

reports_data = {
    1: {
        'rep_no': '01',
        'sub_date': '31/07/2026',
        'date_from': '20/07/2026',
        'date_to': '31/07/2026',
        'about_progress': [
            "The first stage of the BreachShield project focused on establishing the core multi-tier system architecture and engineering the foundational components necessary for privacy-preserving breach detection and out-of-band identity verification. The primary objective during this phase was to build a functional prototype capable of accepting breach lookup queries, enforcing strict authentication, and managing multi-service communication between web interfaces and background OSINT services.",
            "The system architecture was conceptualized as a modular four-tier monorepo consisting of the React 19 web dashboard frontend, the Node.js Express 5 API Gateway, the Python FastAPI OSINT scraper engine, and a native Android Kotlin SMS Gateway companion application. Each tier was designed with clean interface boundaries to ensure high scalability, modularity, and strict access controls.",
            "Significant progress was achieved in developing the backend authentication and security services. The team built the dual-channel One-Time Password (OTP) verification engine (/api/auth/send-otp, /api/auth/verify-otp) supporting both email delivery and mobile SMS relay. Cryptographic security was established using bcrypt hashing for all stored OTPs, a 30-second cooldown timer, a 5-attempt brute-force lockout threshold, and cryptographically signed 1-hour JWT session tokens containing target bindings.",
            "On the data tier, a robust storage model was designed using MySQL 8.0 for user management, authentication sessions, and email_otps tracking, coupled with an automatic JSON-file fallback storage layer. This ensures seamless local development and resilience against database connection interruptions.",
            "The frontend interface was scaffolded as a responsive React 19 single-page application featuring a Cyberpunk-inspired HUD terminal design, dynamic typewriter text streaming, and real-time threat exposure gauges. Strict route gating was implemented to enforce HTTP 403 Forbidden on unauthenticated /api/search requests, preventing unauthorized breach data harvesting.",
            "By the end of Stage 1, BreachShield achieved a fully integrated foundation where users can initiate OTP requests via email or mobile, verify tokens securely, and access guarded search interfaces. This established the foundational security and communication pipelines, creating a solid base for integrating zero-knowledge privacy algorithms and live OSINT scrapers in the subsequent stage."
        ],
        'problems_overcomes': [
            "During Stage 1, several technical challenges emerged involving out-of-band communication, transport authentication, and cross-origin security.",
            "The first major issue involved mobile carrier spam filtering. During initial SMS OTP dispatch testing, outbound messages were silently blocked by Indian telecom DLT (Distributed Ledger Technology) spam filters. The team investigated carrier transmission requirements and resolved this by standardizing the SMS payload into an approved transactional template format.",
            "Another challenge arose with email OTP dispatch via Nodemailer and Gmail SMTP. Developer accounts frequently encountered authentication failures caused by invisible trailing whitespace characters in generated app passwords. The team resolved this by implementing automated regex trimming and sanitization of environment credentials prior to transport initialization.",
            "In the data layer, sporadic connection drops during local testing caused unhandled promise rejections in the backend API. The team engineered an automatic fallback storage abstraction that detects database pool availability and routes persistence seamlessly to local JSON stores without modifying controller logic.",
            "Additionally, cross-origin communication between the React frontend (running on port 3000) and the Node.js backend (on port 5000) led to CORS policy rejections during preflight OPTIONS requests. The team configured centralized CORS middleware with explicit origin whitelisting, allowed headers, and credential support.",
            "Addressing these early integration hurdles fortified the platform's reliability and established a rock-solid foundation for Stage 1."
        ],
        'work_schedule': [
            "Week 1 (July 20 – July 25, 2026): Multi-tier monorepo structure was established across frontend, backend, scraper, and android-gateway. Backend Express 5 server initialization and MySQL 8.0 schema design for users and OTP sessions were completed. Automatic JSON file fallback storage was integrated for development resilience.",
            "Initial system architecture and data flow diagrams were drafted. Core dependencies and libraries were installed and verified for compatibility across Node.js, Python, and React environments.",
            "Week 2 (July 26 – July 31, 2026): Dual-channel OTP authentication engine (send-otp and verify-otp) was implemented with bcrypt hashing and 30-second cooldown timers. Nodemailer Gmail SMTP integration was completed and tested.",
            "Initial WebSocket gateway (/ws/gateway) was scaffolded for Android mobile relays. React 19 Cyberpunk HUD terminal interface was created with strict 403 route gating on search endpoints.",
            "Backend connectivity tests ensured secure token exchange between UI and API routes. Basic session handling and navigation structure were implemented."
        ],
        'final_outcomes': (
            "By the end of Stage 1, BreachShield had evolved into a functional and secure architectural prototype demonstrating the system's core capabilities. "
            "The multi-tier architecture was successfully integrated, enabling seamless communication between the React web console, Express API gateway, and background services. "
            "The system was able to: Enforce strict OTP verification prior to search operations, eliminating unauthenticated breach querying  "
            "Deliver secure OTP tokens through both Gmail SMTP email and WebSocket-relayed mobile SMS  "
            "Protect OTP validation against brute-force attacks using bcrypt hashing, 30s resend cooldowns, and 5-attempt lockouts  "
            "Provide a responsive Cyberpunk HUD terminal user interface with typewriter feedback and live status gauges  "
            "Maintain dual persistence through MySQL 8.0 and resilient local JSON fallback stores  "
            "Demonstrate clean frontend-backend communication with 100% passes on early security challenge tests."
        ),
        'references': [
            "[1] P. Grassi, M. Garcia, and J. Fenton, \"Digital Identity Guidelines: Authentication and Lifecycle Management,\" NIST Special Publication 800-63B, National Institute of Standards and Technology, Gaithersburg, MD, 2020.",
            "[2] OWASP Foundation, \"Authentication Cheat Sheet,\" OWASP Cheat Sheet Series, 2023. Available: https://cheatsheetseries.owasp.org/.",
            "[3] E. Rescorla, \"The Transport Layer Security (TLS) Protocol Version 1.3,\" RFC 8446, Internet Engineering Task Force (IETF), 2018.",
            "[4] M. Conti, A. Dehghantanha, K. Franke, and S. Watson, \"Internet of Things Security and Forensics: Challenges and Opportunities,\" Future Generation Computer Systems, vol. 78, pp. 544–546, 2018.",
            "[5] A. M. Antonopoulos, \"Mastering Bitcoin: Programming the Open Blockchain,\" 2nd ed., O'Reilly Media, Sebastopol, CA, 2017.",
            "[6] RFC 7519, \"JSON Web Token (JWT),\" Internet Engineering Task Force (IETF), 2015. Available: https://datatracker.ietf.org/doc/html/rfc7519."
        ]
    },
    2: {
        'rep_no': '02',
        'sub_date': '15/08/2026',
        'date_from': '01/08/2026',
        'date_to': '15/08/2026',
        'about_progress': [
            "The second stage of the BreachShield project focused on building the privacy-preserving search protocol and implementing the multi-source OSINT aggregation architecture. The primary objective during this period was to ensure that users could search for compromised credentials without ever exposing their raw email addresses or phone numbers to intermediate servers, while aggregating real-time threat intelligence from diverse public and dark web repositories.",
            "A foundational breakthrough was the engineering of the client-side k-anonymity search mechanism. Utilizing the browser's native Web Crypto API (window.crypto.subtle), the user's search query is hashed with SHA-256 locally. Only the first 5 hexadecimal characters of the hash prefix (e.g., e3b0c) are sent to the backend /api/k-anonymity/range/:prefix endpoint. The server responds with all matching hash suffixes in the database, allowing full hash comparison to occur strictly in the client's memory. This guarantees zero server-side exposure of user identifiers.",
            "Concurrently, the team developed the live OSINT scraper microservice in Python using FastAPI and the Telethon MTProto client library. This service connects directly to Telegram threat intelligence channels, monitoring public breach drops, credential dumps, and dark web leak notifications. Concurrency controls using an asyncio.Lock mutex were implemented to serialize MTProto requests, backed by an automated demo_info fallback mechanism during bot rate limits.",
            "To seamlessly combine multiple intelligence channels, the team designed and implemented an extensible object-oriented BreachSource adapter registry. Concrete adapters were developed for: (1) a local indexed catalog of 1,034+ breach records, (2) the live Telegram OSINT scraper, and (3) the Hudson Rock Infostealer intelligence feed. All sources are queried simultaneously via Promise.allSettled, guaranteeing that a failure or timeout in one source never blocks or degrades user queries.",
            "Furthermore, the Android companion application was advanced in native Kotlin. Featuring a modern Liquid Glass UI, a persistent foreground service (GatewayForegroundService.kt), and an intelligent WebSocket manager with exponential backoff and jitter (WebSocketManager.kt), the app securely maintains connection with the backend gateway to relay real-time SMS OTPs and execute remote administrative commands.",
            "By the end of Stage 2, BreachShield possessed an end-to-end privacy-preserving search pipeline supported by live multi-source OSINT harvesting and an active Android gateway relay, setting the stage for machine learning threat scoring and blockchain auditing."
        ],
        'problems_overcomes': [
            "During Stage 2, integration hurdles centered around asynchronous client concurrency, cellular network volatility, and cross-source data normalization.",
            "The primary issue occurred within the Python scraper service: when multiple concurrent search requests hit the Telethon client simultaneously, SQLite threw sqlite3.OperationalError: database is locked. The team resolved this by wrapping all Telethon client invocations in a centralized asyncio.Lock mutex with an 8-second timeout, ensuring safe sequential access to the MTProto session database without blocking the FastAPI event loop.",
            "Another challenge was WebSocket disconnections in the Android gateway application when the mobile device switched between Wi-Fi and 4G/5G cellular networks. The team resolved this by implementing an exponential backoff algorithm with randomized jitter (1s to 30s) and a persistent 5-second HTTP polling fallback (/api/gateway/pending/:deviceId) to guarantee zero lost SMS dispatch requests.",
            "Additionally, search token discrepancies arose when searching Indian mobile numbers due to inconsistent input formatting (e.g., 10-digit numbers vs. numbers with +91 country code). The team implemented an automated normalization preprocessor in searchService.js that standardizes all phone queries into canonical E.164 format.",
            "Finally, handling slow external threat feeds was resolved by transitioning from synchronous batching to non-blocking Promise.allSettled execution, ensuring fast sub-second client response times.",
            "These solutions significantly elevated the robustness and scalability of the multi-source aggregation layer."
        ],
        'work_schedule': [
            "Week 3 (August 01 – August 07, 2026): Client-side k-anonymity prefix hashing was developed using the Web Crypto API. Server-side range lookup endpoint (/api/k-anonymity/range/:prefix) was implemented and tested. Hash suffix matching in React memory was verified.",
            "Python FastAPI scraper microservice was built with Telethon MTProto client integration for Telegram channel monitoring. Mutex locking with asyncio.Lock was added to resolve database lock errors and prevent session corruption.",
            "Week 4 (August 08 – August 15, 2026): Pluggable BreachSource adapter registry was architected and implemented with Promise.allSettled concurrency. Adapters for local breach catalog, Telegram OSINT scraper, and Hudson Rock were integrated.",
            "Android Kotlin companion app was developed with Liquid Glass UI, persistent foreground service, and WebSocketManager.kt backoff reconnects.",
            "Phone number canonicalization preprocessor was deployed. Terminal HUD interface was updated with multi-source aggregated intelligence streams."
        ],
        'final_outcomes': (
            "By the end of Stage 2, BreachShield had achieved a mature privacy-preserving search and multi-source intelligence gathering system. "
            "The multi-source architecture was successfully integrated, enabling seamless communication between the React web dashboard, Express API gateway, Python scraper, and Android mobile gateway. "
            "The system was able to: Perform zero-knowledge credential breach checks using SHA-256 k-anonymity prefix matching, ensuring search queries remain completely private  "
            "Scrape and extract real-time threat intelligence from Telegram leak channels via Telethon with full mutex concurrency protection  "
            "Query multiple heterogeneous threat sources (local catalog, Telegram, Hudson Rock) concurrently with zero single-point-of-failure bottlenecks  "
            "Relay OTP SMS verification messages smoothly through physical Android devices via persistent WebSocket connections with exponential backoff  "
            "Deliver full, sanitized breach intelligence and exposure analytics to the React HUD terminal in sub-second response times."
        ),
        'references': [
            "[1] L. Sweeney, \"k-anonymity: A model for protecting privacy,\" International Journal of Uncertainty, Fuzziness and Knowledge-Based Systems, vol. 10, no. 5, pp. 557–570, 2002.",
            "[2] P. Mell and T. Grance, \"The NIST Definition of Cloud Computing,\" NIST Special Publication 800-145, 2011.",
            "[3] Telethon Documentation, \"Telegram MTProto API Client Library for Python,\" Available: https://docs.telethon.dev/.",
            "[4] A. Aljofey, Q. Jiang, Q. Qu, M. Huang, and J.-P. Niyigena, \"An effective phishing detection model based on character level CNN from URL,\" Electronics, vol. 9, no. 9, p. 1514, 2020.",
            "[5] OWASP Foundation, \"Web Security Testing Guide (WSTG) v4.2,\" 2020. Available: https://owasp.org/www-project-web-security-testing-guide/.",
            "[6] B. Laurie and E. Kasper, \"Certificate Transparency,\" ACM Queue, vol. 12, no. 8, pp. 10–19, 2014."
        ]
    },
    3: {
        'rep_no': '03',
        'sub_date': '04/09/2026',
        'date_from': '16/08/2026',
        'date_to': '04/09/2026',
        'about_progress': [
            "The third stage of the BreachShield project focused on machine learning threat classification, blockchain-anchored tamper-proof audit logging, and comprehensive full-stack security hardening. The primary objective was to replace heuristic threat placeholders with real, empirical AI models and build a verifiable cryptographic audit trail for all breach search transactions.",
            "In the AI threat analysis module, the team curated a verified benchmark dataset of 1,034 breach records sourced from the HaveIBeenPwned catalog. A rigorous machine learning pipeline was constructed to classify breach severity based on 7 feature dimensions, including compromised account volume, credential sensitivity (passwords, PII, financial info), and breach recency. Four distinct model architectures were trained and evaluated across a stratified 80/20 train/test split: XGBoost achieved 89.2% accuracy and was selected for production classification over deep learning models (CNN: 66.2%, RNN: 44.6%, Transformer: 48.4%).",
            "Concurrently, the team engineered a decentralized, tamper-proof audit logging infrastructure. Using Solidity, the smart contract AnchorRegistry.sol was authored with anchorRoot(bytes32) and verifyRoot(bytes32) functions. In the Node.js backend, merkleBatcher.js was built to buffer search events, compute SHA-256 leaf hashes, build balanced Merkle trees, and anchor root hashes on-chain using ethers.js. An automated Hardhat local blockchain lifecycle with startup deployment hooks (ensureContractDeployed()) and a dead-letter queue (DLQ) retry mechanism was implemented for flawless test reproducibility.",
            "A rigorous, full-stack security and penetration testing audit was conducted across the monorepo. The team resolved 5 critical security concerns and 5 logical bugs: unauthenticated endpoints were locked down, JWT signing keys were rotated to 128-character crypto-random hex values, OTP generation was migrated from Math.random() to crypto.randomInt(), and server logging was sanitized to prevent secret leaks.",
            "To demonstrate system health, the automated test suite was expanded to 94 backend tests (covering integration, boundary, and adversarial challenge suites) and 26 frontend component tests, achieving a 100% clean pass rate. Additionally, a real-time darkweb watchlist stream using Server-Sent Events (SSE) and an offline Ollama LLM integration were integrated for local threat explanation.",
            "By the end of Stage 3, BreachShield represented a hardened, production-grade cybersecurity platform that unites AI-driven risk scoring, zero-knowledge privacy, verifiable blockchain audit trails, and multi-platform mobile relays."
        ],
        'problems_overcomes': [
            "During Stage 3, challenges involved blockchain network stability, deep learning model suitability on tabular data, and state persistence during server restarts.",
            "The first major hurdle involved public testnet deployment. Initial tests anchoring Merkle roots onto the Polygon Amoy public testnet encountered severe RPC rate-limiting, unpredictable gas fee spikes, and frequent transaction timeouts that disrupted automated CI runs. The team solved this by engineering an automated local Hardhat node runner with auto-deployment hooks, ensuring deterministic, instant, zero-cost anchoring during demonstrations while documenting public testnet migration readiness.",
            "Another critical challenge was the poor performance of deep learning models (CNN, RNN, Transformer), which achieved only 44%–66% accuracy on breach metadata due to the tabular, non-sequential nature of breach catalog features. The team overcame this by pivoting to XGBoost, optimizing hyperparameter trees, and achieving a production-ready 89.2% accuracy.",
            "In the audit pipeline, testing revealed that restarting the backend server while a batch was pending caused Merkle proofs to be lost upon recovery, resulting in proof verification failures. The team fixed this by persisting generated proofs in the batch record before writing to history.",
            "Furthermore, an adversarial security audit identified that POST /api/search catch blocks were swallowing 500 errors and presenting a false 'clean scan' UI. The team refactored error handling across all controllers to return authentic HTTP status codes.",
            "Overcoming these real-world obstacles ensured BreachShield met the highest standards of architectural integrity and academic rigor."
        ],
        'work_schedule': [
            "Week 5 (August 16 – August 24, 2026): Sourced and curated the 1,034-record HaveIBeenPwned benchmark dataset. Built feature extraction pipeline and trained four ML models (XGBoost, CNN, RNN, Transformer). Evaluated test set accuracy and selected XGBoost (89.2% accuracy).",
            "Developed Solidity AnchorRegistry.sol smart contract and built merkleBatcher.js with cryptographic leaf hashing.",
            "Week 6–7 (August 25 – September 04, 2026): Automated local Hardhat blockchain lifecycle with boot deployer hook and dead-letter queue (DLQ) retry mechanisms. Conducted comprehensive security penetration audit, rotating JWT keys and securing OTP generation.",
            "Expanded backend test suite to 94 tests and frontend to 26 tests (100% passing). Integrated darkweb SSE stream and offline Ollama LLM connector.",
            "Finalized Phase 2 presentation slides and technical documentation. Rehearsed live multi-tier demonstrations."
        ],
        'final_outcomes': (
            "By the end of Stage 3, BreachShield had achieved a fully realized, verified, and hardened cybersecurity architecture. "
            "The multi-module platform was successfully integrated, enabling seamless communication between the React web console, Express API gateway, Python OSINT scraper, Android companion app, and Ethereum/Hardhat blockchain. "
            "The system was able to: Accurately assess breach severity in real time using a trained XGBoost machine learning model achieving 89.2% accuracy  "
            "Maintain a tamper-evident audit trail of all breach queries by constructing Merkle trees and anchoring root hashes on a Solidity smart contract  "
            "Pass all 94 backend tests and 26 frontend component tests with 0 errors, withstanding adversarial fuzzing and token tampering  "
            "Provide real-time darkweb intelligence streaming via Server-Sent Events (SSE) and explain discovered threats locally using offline Ollama LLMs  "
            "Provide complete administrative controls via the native Android companion app and React HUD dashboard."
        ),
        'references': [
            "[1] T. Chen and C. Guestrin, \"XGBoost: A scalable tree boosting system,\" in Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 2016, pp. 785–794.",
            "[2] R. C. Merkle, \"A Digital Signature Based on a Conventional Encryption Function,\" in Advances in Cryptology — CRYPTO '87, Lecture Notes in Computer Science, vol. 293, Springer, Berlin, Heidelberg, 1988, pp. 369–378.",
            "[3] S. Nakamoto, \"Bitcoin: A Peer-to-Peer Electronic Cash System,\" 2008. Available: https://bitcoin.org/bitcoin.pdf.",
            "[4] OWASP Foundation, \"Testing JSON Web Tokens,\" OWASP Web Security Testing Guide (WSTG), 2024. Available: https://owasp.org/www-project-web-security-testing-guide/.",
            "[5] A. Vaswani et al., \"Attention Is All You Need,\" in Advances in Neural Information Processing Systems (NeurIPS 2017), vol. 30, 2017.",
            "[6] Ethereum Foundation, \"Solidity Documentation and Smart Contract Guidelines,\" Available: https://docs.soliditylang.org/."
        ]
    }
}

target_title = 'BreachShield: AI Powered OSINT Breach Detection & Dark Web Protection'

for num in [1, 2, 3]:
    info = reports_data[num]
    
    # 1. Fresh load from pristine template
    doc = docx.Document(template_path)
    
    # 2. Update Table 0 (Metadata)
    t0 = doc.tables[0]
    t0.rows[0].cells[0].text = 'PROJECT WORK PHASE-2 PROGRESS REPORT'
    t0.rows[1].cells[2].text = '05'
    t0.rows[2].cells[2].text = 'Prof. Sushma T.M.'
    t0.rows[3].cells[0].text = 'Project Title'
    t0.rows[3].cells[2].text = target_title
    t0.rows[4].cells[2].text = info['rep_no']
    t0.rows[5].cells[2].text = info['sub_date']
    t0.rows[6].cells[2].text = f"From: {info['date_from']}"
    t0.rows[6].cells[4].text = f"To: {info['date_to']}"
    
    # Students
    students = [
        ('1', 'Arvind D H', '1AY24IS400'),
        ('2', 'Girishkumar', '1AY24IS404'),
        ('3', 'Manoj', '1AY24IS406'),
        ('4', 'Premkumar', '1AY24IS407')
    ]
    for idx, (s_no, s_name, s_usn) in enumerate(students):
        row = t0.rows[9 + idx]
        row.cells[0].text = s_no
        row.cells[1].text = s_name
        row.cells[3].text = s_usn
        
    # 3. Update Table 1 (Signatures)
    t1 = doc.tables[1]
    t1.rows[0].cells[0].text = 'Prof. Sushma T.M.'
    t1.rows[0].cells[1].text = 'Dr. Pankaj Kumar'
    t1.rows[0].cells[2].text = 'Prof. Yogesh N'
    t1.rows[0].cells[3].text = 'HOD-ISE'
    for c in t1.rows[1].cells:
        c.text = '(Signature with Date)'
        
    # 4. Populate Paragraphs
    for i, p_text in enumerate(info['about_progress']):
        doc.paragraphs[9 + i].text = p_text
        doc.paragraphs[9 + i].paragraph_format.line_spacing = 1.15
        doc.paragraphs[9 + i].paragraph_format.space_after = Pt(2.5)
        
    for i, p_text in enumerate(info['problems_overcomes']):
        doc.paragraphs[17 + i].text = p_text
        doc.paragraphs[17 + i].paragraph_format.line_spacing = 1.15
        doc.paragraphs[17 + i].paragraph_format.space_after = Pt(2.5)
        
    for i, p_text in enumerate(info['work_schedule']):
        doc.paragraphs[27 + i].text = p_text
        doc.paragraphs[27 + i].paragraph_format.line_spacing = 1.15
        doc.paragraphs[27 + i].paragraph_format.space_after = Pt(2.5)
        
    doc.paragraphs[34].text = info['final_outcomes']
    doc.paragraphs[34].paragraph_format.line_spacing = 1.15
    doc.paragraphs[34].paragraph_format.space_after = Pt(2.5)
    
    doc.paragraphs[40].text = info['references'][0]
    doc.paragraphs[41].text = info['references'][1]
    doc.paragraphs[42].text = info['references'][2]
    doc.paragraphs[43].text = ''  # empty leftover line
    doc.paragraphs[44].text = info['references'][3]
    doc.paragraphs[45].text = info['references'][4]
    doc.paragraphs[46].text = info['references'][5]
    for p_idx in [40, 41, 42, 44, 45, 46]:
        doc.paragraphs[p_idx].paragraph_format.line_spacing = 1.15
        doc.paragraphs[p_idx].paragraph_format.space_after = Pt(2.5)
        
    # 5. Remove trailing empty paragraphs after Table 1
    body = doc._body._element
    tables = body.xpath('.//w:tbl')
    last_tbl = tables[-1]
    last_idx = list(body).index(last_tbl)
    
    to_remove = []
    for i in range(last_idx + 1, len(body)):
        el = body[i]
        if el.tag.endswith('p'):
            to_remove.append(el)
            
    for el in to_remove:
        body.remove(el)
    print(f'Report {num}: Removed {len(to_remove)} trailing empty paragraphs after Table 1')
    
    # 6. Apply native double-line pgBorders to all sections using oxml
    for section in doc.sections:
        sectPr = section._sectPr
        existing = sectPr.find(docx.oxml.ns.qn('w:pgBorders'))
        if existing is not None:
            sectPr.remove(existing)
            
        pg_borders = parse_xml(
            r'<w:pgBorders %s w:offsetFrom="page">'
            r'  <w:top w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
            r'  <w:left w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
            r'  <w:bottom w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
            r'  <w:right w:val="double" w:sz="12" w:space="24" w:color="000000"/>'
            r'</w:pgBorders>' % nsdecls('w')
        )
        sectPr.append(pg_borders)
        
    # 7. Save clean docx files
    final_docx = os.path.join(academic_dir, f'BreachShield_Progress_Report_{num}.docx')
    draft_docx = os.path.join(academic_dir, f'BreachShield_Progress_Report_{num}_DRAFT.docx')
    
    doc.save(final_docx)
    try:
        doc.save(draft_docx)
    except PermissionError:
        pass
    print(f'Report {num}: Successfully generated {final_docx}')

print('All 3 reports successfully built!')
