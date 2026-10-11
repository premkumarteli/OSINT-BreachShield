import pptx

prs = pptx.Presentation(r'C:\Users\prem\OSINT-breach-Finder-main\docs\Phase2\Final_presentation2 format 2026-27_BLANK_TEMPLATE.pptx')
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for sp in slide.shapes:
        txt = sp.text.strip().replace('\n', ' [NL] ') if sp.has_text_frame else ''
        print(f"  Shape [{sp.name}] ID={sp.shape_id} pos=({sp.left.inches:.2f}, {sp.top.inches:.2f}, {sp.width.inches:.2f}, {sp.height.inches:.2f})")
        if txt:
            print(f"     text: \"{txt[:120]}\"")
