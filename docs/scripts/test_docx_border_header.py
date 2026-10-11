import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(1.25)
sec.right_margin = Inches(1.0)
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)

# 1. Add VTU thinThickSmallGap double page border
border_xml = parse_xml(r'''
    <w:pgBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:offsetFrom="page">
        <w:top w:val="thinThickSmallGap" w:sz="24" w:space="24" w:color="auto"/>
        <w:left w:val="thinThickSmallGap" w:sz="24" w:space="31" w:color="auto"/>
        <w:bottom w:val="thickThinSmallGap" w:sz="24" w:space="24" w:color="auto"/>
        <w:right w:val="thickThinSmallGap" w:sz="24" w:space="24" w:color="auto"/>
    </w:pgBorders>
''')
sec._sectPr.append(border_xml)

# 2. Add Running Header
header = sec.header
p_head = header.paragraphs[0]
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_head = p_head.add_run("BreachShield: AI Powered OSINT Breach Detection & Dark Web Protection")
r_head.font.name = "Times New Roman"
r_head.font.size = Pt(9.5)
r_head.font.italic = True
r_head.font.color.rgb = RGBColor(100, 100, 100)

# Add bottom border to header paragraph
pPr = p_head._p.get_or_add_pPr()
pBdr = parse_xml(r'''
    <w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:bottom w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/>
    </w:pBdr>
''')
pPr.append(pBdr)

# 3. Add Running Footer
footer = sec.footer
p_foot = footer.paragraphs[0]
p_foot.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Add top border to footer paragraph
pPr_f = p_foot._p.get_or_add_pPr()
pBdr_f = parse_xml(r'''
    <w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
        <w:top w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/>
    </w:pBdr>
''')
pPr_f.append(pBdr_f)

r_foot_l = p_foot.add_run("Dept. of ISE, AcIT                                                               2026 - 2027\t")
r_foot_l.font.name = "Times New Roman"
r_foot_l.font.size = Pt(9.5)
r_foot_l.font.bold = True
r_foot_l.font.color.rgb = RGBColor(80, 80, 80)

# Add dynamic PAGE number field
fldSimple = OxmlElement('w:fldSimple')
fldSimple.set(qn('w:instr'), 'PAGE')
p_foot._p.append(fldSimple)

p = doc.add_paragraph("This is a test paragraph to verify borders, headers, and footers.")
test_out = r"C:\Users\prem\OSINT-breach-Finder-main\docs\scripts\test_border_out.docx"
doc.save(test_out)
print("Saved test document successfully!")
